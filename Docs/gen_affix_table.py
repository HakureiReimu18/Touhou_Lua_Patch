# -*- coding: utf-8 -*-
# 词缀一览表同步器：用 XML/文本文件的内容【原地更新】Docs/词缀一览表.xlsx。
# 设计约束：
#   1. 只改单元格的 value，绝不触碰任何格式（字体/底色/边框/列宽/行高/合并全部保留）；
#   2. 按 ID 列（表头"ID"所在列）匹配行，不改动行顺序；
#   3. XML 里新增的词缀追加到表尾（样式复制自最后一行数据），XML 里已删除的词缀不自动删除（只提示）；
#   4. 表格不存在时才全新创建。保存走临时文件 + 原子替换，失败不会损坏原表。
# 用法：python gen_affix_table.py（文件正被 Excel 打开时无法保存，会给出提示）
import os
import re
import shutil
import tempfile
import xml.etree.ElementTree as ET

MOD_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(MOD_ROOT, "Docs", "词缀一览表.xlsx")

TIER_ORDER = ["Broken", "Normal", "Rare", "Epic", "Legendary", "Special"]
TIER_CN = {"Broken": "破损", "Normal": "普通", "Rare": "稀有", "Epic": "史诗", "Legendary": "传说", "Special": "特殊"}
TIER_FILL = {
    "Broken": "FF808080", "Normal": "FFFFFFFF", "Rare": "FF4A90FF",
    "Epic": "FFC040FF", "Legendary": "FFFF8C00", "Special": "FFFF4040",
}
APPLICABLE_CN = {
    "weapon": "武器", "meleeweapon": "近战武器", "rangedweapon": "远程武器",
    "tool": "工具", "repairtool": "修复工具", "medical": "医疗", "armor": "护甲", "wearable": "穿戴物", "all": "全部",
}
HEADERS = ["稀有度", "名称", "ID", "适用范围", "具体效果"]


def strip_rich(s):
    return re.sub(r"‖[^‖]*‖", "", s or "").strip()


def parse_affixes():
    """从 XML + 中文本地化解析全部词缀 → {id: (稀有度文本, 名称, 适用范围, 效果)}"""
    affixes = {}
    tree = ET.parse(os.path.join(MOD_ROOT, "Items", "Affixes.xml"))
    for el in tree.getroot().iter("Affix"):
        aid = el.get("identifier")
        affixes[aid] = {
            "tier": el.get("tier", "Normal"),
            "applicable": el.get("applicable", "all"),
            "desc": el.get("desc", ""),
            "descarmor": el.get("descarmor", ""),
            "name": "",
            "order": len(affixes),
        }

    loc = ET.parse(os.path.join(MOD_ROOT, "Text", "Affixes_SChinese.xml"))
    for el in loc.getroot():
        m = re.match(r"affix\.(prefix|desc|descarmor)\.(.+)", el.tag)
        if not m:
            continue
        kind, aid = m.group(1), m.group(2)
        if aid not in affixes:
            continue
        text = el.text or ""
        if kind == "prefix":
            affixes[aid]["name"] = strip_rich(text)
        elif kind == "desc":
            affixes[aid]["desc"] = text.strip()
        elif kind == "descarmor":
            affixes[aid]["descarmor"] = text.strip()

    result = {}
    for aid, a in affixes.items():
        effect = a["desc"]
        if a["descarmor"] and a["descarmor"] != a["desc"]:
            effect = (effect + "　【护甲】" + a["descarmor"]) if effect else "【护甲】" + a["descarmor"]
        app_cn = "、".join(APPLICABLE_CN.get(t.strip(), t.strip()) for t in a["applicable"].split(","))
        result[aid] = (f"{TIER_CN[a['tier']]}（{a['tier']}）", a["name"] or aid, app_cn, effect)
    return result


def find_columns(ws):
    """按表头文字定位五列（容忍用户调整列顺序/插入列）"""
    cols = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=1, column=col).value
        if v is None:
            continue
        name = str(v).strip()
        if name in HEADERS and name not in cols:
            cols[name] = col
    return cols


def main():
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    affixes = parse_affixes()

    if not os.path.exists(OUT_PATH):
        # 表格不存在：全新创建（带默认格式）
        wb = Workbook()
        ws = wb.active
        ws.title = "词缀一览"
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.alignment = Alignment(horizontal="center", vertical="center")
        rows = sorted(affixes.items(), key=lambda kv: (TIER_ORDER.index(kv[1][0].split("（")[1][:-1]), kv[1][2], kv[0]))
        for aid, (tier_text, name, app_cn, effect) in rows:
            ws.append([tier_text, name, aid, app_cn, effect])
            cell = ws.cell(row=ws.max_row, column=1)
            tier = tier_text.split("（")[1][:-1]
            cell.fill = PatternFill(start_color=TIER_FILL[tier], end_color=TIER_FILL[tier], fill_type="solid")
            cell.font = Font(color="FF000000", bold=True)
        for i, w in enumerate([16, 14, 18, 18, 80], 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="center", wrap_text=(c.column == 5))
        ws.freeze_panes = "A2"
        wb.save(OUT_PATH)
        print(f"created: {OUT_PATH} ({len(rows)} affixes)")
        return

    # 已存在：原地更新，格式零改动
    wb = load_workbook(OUT_PATH)
    ws = wb.active
    cols = find_columns(ws)
    missing = [h for h in HEADERS if h not in cols]
    if missing:
        raise SystemExit(f"表头缺列 {missing}，请先恢复表头（第 1 行）；未做任何修改")

    id_col = cols["ID"]
    row_of = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=id_col).value
        if v is not None and str(v).strip():
            row_of[str(v).strip()] = r

    updated = 0
    appended = []
    for aid, (tier_text, name, app_cn, effect) in affixes.items():
        values = {"稀有度": tier_text, "名称": name, "适用范围": app_cn, "具体效果": effect}
        if aid in row_of:
            r = row_of[aid]
            for header, val in values.items():
                cell = ws.cell(row=r, column=cols[header])
                if cell.value != val:
                    cell.value = val  # 只改 value，样式不动
                    updated += 1
        else:
            appended.append((aid, tier_text, name, app_cn, effect))

    # 新增词缀：追加到表尾，样式复制最后一行数据行
    for aid, tier_text, name, app_cn, effect in appended:
        r = ws.max_row + 1
        src = max(2, r - 1)
        for header, val in zip(HEADERS, [tier_text, name, aid, app_cn, effect]):
            cell = ws.cell(row=r, column=cols[header])
            cell.value = val
            style_src = ws.cell(row=src, column=cols[header])
            if style_src.has_style:
                cell._style = style_src._style  # 复制格式（不引入新格式）

    stale = [aid for aid in row_of if aid not in affixes]
    if stale:
        print(f"提示：表格里存在 XML 中已删除的词缀行了，未自动删除：{', '.join(stale)}")

    # 临时文件 + 原子替换：保存失败不会损坏原表
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(OUT_PATH))
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, OUT_PATH)
    except PermissionError:
        os.remove(tmp)
        raise SystemExit("保存失败：表格正被 Excel 打开，请关闭 Excel 后重试（未做任何修改）")
    except Exception:
        if os.path.exists(tmp):
            os.remove(tmp)
        raise

    print(f"updated: {OUT_PATH}（修改 {updated} 个单元格，追加 {len(appended)} 行{('：' + ', '.join(a[0] for a in appended)) if appended else ''}）")


if __name__ == "__main__":
    main()
